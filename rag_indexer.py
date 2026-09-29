"""Build and query a Haystack + PostgreSQL/pgvector PDF RAG index."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import unquote, urlparse

from haystack import Document
from haystack.components.converters import PyPDFToDocument
from haystack.components.embedders import SentenceTransformersDocumentEmbedder, SentenceTransformersTextEmbedder
from haystack.components.preprocessors import DocumentSplitter
from haystack_integrations.components.retrievers.pgvector import PgvectorEmbeddingRetriever
from haystack.components.writers import DocumentWriter
from haystack_integrations.document_stores.pgvector import PgvectorDocumentStore
from haystack.utils import Secret
from openpyxl import load_workbook


URL_PATTERN = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
METADATA_KEY_PATTERN = re.compile(r"[^a-zA-Z0-9_]+")
DEFAULT_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
CANONICAL_METADATA_FIELDS = {
    "fm_id": "failure_id",
    "failure_id": "failure_id",
    "fm_class": "failure_class",
    "category": "failure_class",
    "failure_mode_fm": "failure_mode",
    "specific_failure_mode_threat_anomaly_name": "failure_mode",
    "potential_causes_mechanism_s": "causes",
    "potential_causes_mechanism_s_of_the_fm": "causes",
    "causes": "causes",
    "detailed_description_technical_mechanism": "technical_description",
    "potential_technical_effects_on_the_asset": "technical_effects",
    "potential_technical_effects_of_the_fm": "technical_effects",
    "potential_consequences": "consequences",
    "potential_consequences_of_the_fm": "consequences",
    "severity_1_10": "severity",
    "probability_1_10": "probability",
    "rpn_s_x_p": "rpn",
    "detective_controls": "detective_controls",
    "preventive_controls": "preventive_controls",
    "responsibility": "responsibility",
    "source": "source",
    "link": "source_link",
}


def url_filename(url: str) -> str:
    name = Path(unquote(urlparse(url).path)).name
    return name.rstrip(" .") or "download"


def metadata_key(value: str) -> str:
    normalized = METADATA_KEY_PATTERN.sub("_", value).strip("_").lower()
    return normalized or "field"


def excel_records(data_dir: Path) -> list[dict[str, object]]:
    """Read every linked row and retain all populated Excel fields as metadata."""
    records: list[dict[str, object]] = []
    for workbook_path in sorted(data_dir.glob("*.xlsx")):
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows()
                headers: list[str] | None = None
                for row_number, row in enumerate(rows, start=1):
                    values = [cell.value for cell in row]
                    if headers is None:
                        if any(str(value).strip().lower() == "link" for value in values if value is not None):
                            headers = [str(value).strip() if value is not None else "" for value in values]
                        continue
                    metadata: dict[str, object] = {}
                    links: list[str] = []
                    for index, value in enumerate(values):
                        key = headers[index] if index < len(headers) else ""
                        if not key or key.lower() == "column guide":
                            continue
                        if value is not None and str(value).strip():
                            metadata[key] = value
                            if key.lower() == "link":
                                links.extend(URL_PATTERN.findall(str(value)))
                    for link in links:
                        records.append({
                            "url": link.rstrip(".,;:)]}"),
                            "filename": url_filename(link),
                            "source_workbook": workbook_path.name,
                            "source_sheet": worksheet.title,
                            "source_row": row_number,
                            "excel_metadata": metadata,
                        })
        finally:
            workbook.close()
    return records


def metadata_by_filename(data_dir: Path) -> dict[str, dict[str, object]]:
    by_filename: dict[str, dict[str, object]] = {}
    for record in excel_records(data_dir):
        filename = str(record["filename"]).lower()
        by_filename.setdefault(filename, record)
    return by_filename


def document_metadata(pdf_path: Path, lookup: dict[str, dict[str, object]]) -> dict[str, object]:
    record = lookup.get(pdf_path.name.lower())
    if record is None:
        record = lookup.get(pdf_path.stem.lower())
    metadata: dict[str, object] = {
        "file_name": pdf_path.name,
        "file_path": str(pdf_path.resolve()),
    }
    if record:
        metadata.update({key: value for key, value in record.items() if key != "excel_metadata"})
        for key, value in record["excel_metadata"].items():
            normalized_key = metadata_key(key)
            metadata[f"excel_{normalized_key}"] = value
            canonical_key = CANONICAL_METADATA_FIELDS.get(normalized_key)
            if canonical_key:
                metadata[canonical_key] = value
    else:
        metadata["metadata_match"] = False
    return metadata


def build_index(pdf_dir: Path, data_dir: Path, connection_string: str, model: str) -> int:
    pdf_paths = sorted(pdf_dir.glob("*.pdf"))
    if not pdf_paths:
        print(f"No PDF files found in {pdf_dir}", file=sys.stderr)
        return 1
    document_store = PgvectorDocumentStore(
        connection_string=Secret.from_token(connection_string),
        table_name="rag_documents",
        embedding_dimension=384,
        create_extension=True,
        recreate_table=True,
        search_strategy="hnsw",
    )
    converter = PyPDFToDocument()
    splitter = DocumentSplitter(split_by="word", split_length=200, split_overlap=30)
    embedder = SentenceTransformersDocumentEmbedder(model=model)
    embedder.warm_up()
    writer = DocumentWriter(document_store=document_store)
    lookup = metadata_by_filename(data_dir)
    documents: list[Document] = []
    for pdf_path in pdf_paths:
        converted = converter.run(sources=[pdf_path])["documents"]
        source_metadata = document_metadata(pdf_path, lookup)
        chunks = splitter.run(documents=converted)["documents"]
        for document in chunks:
            document.meta.update(source_metadata)
            documents.append(document)
    embedded = embedder.run(documents=documents)["documents"]
    writer.run(documents=embedded)
    print(f"Indexed {len(pdf_paths)} PDF(s), {len(embedded)} chunk(s) into PostgreSQL table rag_documents")
    return 0


def query_index(connection_string: str, question: str, model: str, top_k: int, filter_values: list[str]) -> int:
    document_store = PgvectorDocumentStore(
        connection_string=Secret.from_token(connection_string),
        table_name="rag_documents",
        embedding_dimension=384,
        create_extension=True,
        search_strategy="hnsw",
    )
    embedder = SentenceTransformersTextEmbedder(model=model)
    embedder.warm_up()
    query_embedding = embedder.run(text=question)["embedding"]
    retriever = PgvectorEmbeddingRetriever(document_store=document_store)
    conditions = []
    for item in filter_values:
        key, separator, value = item.partition("=")
        if not separator or not key:
            raise ValueError(f"Invalid filter {item!r}; expected key=value")
        conditions.append({"field": key if key.startswith("meta.") else f"meta.{key}", "operator": "==", "value": value})
    filters = {"operator": "AND", "conditions": conditions} if conditions else None
    results = retriever.run(query_embedding=query_embedding, top_k=top_k, filters=filters)["documents"]
    payload = [{"score": document.score, "content": document.content, "meta": document.meta} for document in results]
    print(json.dumps(payload, indent=2, default=str))
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf-dir", type=Path, default=Path("downloaded"))
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--connection-string", default=os.getenv("PG_CONN_STR"), help="PostgreSQL URL; defaults to PG_CONN_STR")
    parser.add_argument("--model", default=DEFAULT_MODEL)
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("index", help="Index downloaded PDFs")
    query = subparsers.add_parser("query", help="Query the vector index")
    query.add_argument("question")
    query.add_argument("--top-k", type=int, default=5)
    query.add_argument("--filter", action="append", default=[], dest="filter_values", help="Metadata filter as key=value; repeat for more fields")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.connection_string:
        print("Set PG_CONN_STR or pass --connection-string.", file=sys.stderr)
        return 2
    if args.command == "index":
        return build_index(args.pdf_dir, args.data_dir, args.connection_string, args.model)
    return query_index(args.connection_string, args.question, args.model, args.top_k, args.filter_values)


if __name__ == "__main__":
    raise SystemExit(main())